import json
import os
from pathlib import Path
import subprocess
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import Mock, patch

from agent.file_writer import FileWriter, PandocBackend, FileWriteError, pdf_symbol_text
from agent.main import run_turn
from agent.tools import FileTools, TOOLS

BODY = (Path(__file__).parent / "fixtures/file_writer.md").read_text(encoding="utf-8")


class FileWriterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.project = Path(self.temp.name)
        self.scope = patch("agent.file_writer.PROJECT", self.project)
        self.scope.start()
        self.addCleanup(self.scope.stop)
        self.writer = FileWriter()

    def test_real_four_formats_chinese_and_structure(self):
        from docx import Document
        import pdfplumber
        self.assertFalse(self.writer.output_dir.exists())
        for ext in ("txt", "md", "docx", "pdf"):
            with self.subTest(extension=ext):
                result = self.writer.create_file("test." + ext, BODY)
                self.assertTrue(result["success"], result)
                output = Path(result["path"])
                self.assertEqual(output.parent, self.project / "generated_files")
                self.assertGreater(output.stat().st_size, 0)
                if ext in ("txt", "md"):
                    self.assertEqual(output.read_bytes(), BODY.encode("utf-8"))
                elif ext == "docx":
                    doc = Document(output)
                    text = "\n".join(p.text for p in doc.paragraphs)
                    self.assertIn("中文标题", text)
                    self.assertIn("314159", text)
                    self.assertTrue(any(p.style.name.startswith("Heading") for p in doc.paragraphs))
                    self.assertTrue(any(r.bold for p in doc.paragraphs for r in p.runs))
                    self.assertIn("中文", str([[c.text for c in row.cells] for row in doc.tables[0].rows]))
                    self.assertTrue(any(p._p.xpath("./w:pPr/w:numPr") for p in doc.paragraphs))
                else:
                    with pdfplumber.open(output) as pdf:
                        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
                    for expected in ("中文标题", "中文段落", "代码块", "314159", "English Heading", "正常显示"):
                        self.assertIn(expected, text)
        self.assertEqual(len(list(self.writer.output_dir.iterdir())), 4)

    def test_invalid_names_and_extensions(self):
        for name in ("", "../a.md", "..\\a.md", "folder/a.md", "C:\\Desktop\\a.md", "/tmp/a.md",
                     "\\\\server\\a.md", "NUL.txt", "a.md:stream", "a.md ", "a\x00.md"):
            with self.subTest(name=name):
                self.assertEqual(self.writer.create_file(name, BODY)["code"], "invalid_filename")
        for name in ("a.xlsx", "a.pptx", "a.html", "a"):
            self.assertEqual(self.writer.create_file(name, BODY)["code"], "unsupported_type")
        self.assertFalse(self.writer.output_dir.exists())

    def test_pdf_chinese_and_symbols_in_heading_table_and_code(self):
        import pdfplumber
        body = ("# 中文标题 ✅❌\n\n中文✅正确❌错误 English ⚠️ 📝💻📄\n\n"
                "**✅ 加粗**\n\n| 状态 | 说明 |\n| --- | --- |\n| ❌ | 失败 |\n\n"
                "```text\n中文代码✅❌⚠️\n```\n\n`✅行内代码`\n")
        result = self.writer.create_file("symbols.pdf", body)
        self.assertTrue(result["success"], result)
        with pdfplumber.open(result["path"]) as pdf:
            text = "\n".join(p.extract_text() or "" for p in pdf.pages)
            fonts = {c["fontname"] for p in pdf.pages for c in p.chars}
        for symbol in "✅❌⚠📝💻📄":
            self.assertIn(symbol, text)
        self.assertIn("中文代码✅❌⚠", text.replace(" ", ""))
        self.assertIn("✅行内代码", text.replace(" ", ""))
        self.assertTrue(any("SegoeUISymbol" in font for font in fonts), fonts)
        self.assertTrue(any("YaHei" in font for font in fonts), fonts)
        self.assertEqual(len(list(self.writer.output_dir.iterdir())), 1)

    def test_presentation_selectors_only_normalized_for_pdf_symbols(self):
        body = "⚠️✅︎中文\uFE0F 字母A\uFE0F"
        self.assertEqual(pdf_symbol_text(body), "⚠✅中文\uFE0F 字母A\uFE0F")
        for ext in ("txt", "md", "docx"):
            result = self.writer.create_file("unchanged." + ext, body)
            self.assertTrue(result["success"], result)
            if ext == "docx":
                from docx import Document
                actual = "\n".join(p.text for p in Document(result["path"]).paragraphs)
            else:
                actual = Path(result["path"]).read_text(encoding="utf-8")
            self.assertEqual(actual, body)

    def test_unknown_glyph_remains_an_error_with_exact_codepoint(self):
        result = self.writer.create_file("unsupported.pdf", "中文\u0378")
        self.assertFalse(result["success"], result)
        self.assertEqual(result["code"], "missing_glyph")
        self.assertIn("U+0378", result["error"])
        self.assertFalse((self.writer.output_dir / "unsupported.pdf").exists())
        self.assertEqual(list(self.writer.output_dir.iterdir()), [])

    def test_missing_glyph_error_distinguishes_symbols_from_chinese(self):
        stderr = ("[WARNING] Missing character: There is no ✅ (U+2705) (U+2705) in font\n"
                  "[WARNING] Missing character: There is no ❌ (U+274C) (U+274C) in font\n")
        with patch("agent.file_writer.subprocess.run", return_value=subprocess.CompletedProcess([], 0, "", stderr)):
            with self.assertRaises(FileWriteError) as error:
                PandocBackend()._run(["pandoc"], self.project, "body")
        self.assertIn("✅", str(error.exception))
        self.assertIn("❌", str(error.exception))
        self.assertEqual(str(error.exception).count("U+2705"), 1)
        self.assertNotIn("请检查中文字体", str(error.exception))

    def test_empty_and_large_content(self):
        for body in ("", "  ", None, "x" * 80001):
            self.assertFalse(self.writer.create_file("a.md", body)["success"])

    def test_no_overwrite_or_automatic_rename(self):
        self.assertTrue(self.writer.create_file("same.md", "original")["success"])
        self.assertEqual(self.writer.create_file("same.md", "changed")["code"], "file_exists")
        self.assertEqual((self.writer.output_dir / "same.md").read_text(), "original")
        self.assertEqual(len(list(self.writer.output_dir.iterdir())), 1)

    def test_concurrent_creation_is_exclusive(self):
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda text: self.writer.create_file("race.txt", text), ["one", "two"]))
        self.assertEqual(sum(r["success"] for r in results), 1)
        self.assertIn((self.writer.output_dir / "race.txt").read_text(), ("one", "two"))
        self.assertEqual(len(list(self.writer.output_dir.iterdir())), 1)

    def test_directory_creation_and_write_permission_errors(self):
        with patch.object(Path, "mkdir", side_effect=PermissionError()):
            self.assertEqual(self.writer.create_file("a.md", BODY)["code"], "access_denied")
        with patch.object(Path, "mkdir", side_effect=OSError("disk failure")):
            self.assertEqual(self.writer.create_file("a.md", BODY)["code"], "filesystem_failed")
        with patch("agent.file_writer.os.link", side_effect=PermissionError()):
            self.assertEqual(self.writer.create_file("a.md", BODY)["code"], "access_denied")
        self.assertEqual(list(self.writer.output_dir.iterdir()), [])

    def test_output_directory_cannot_be_a_file(self):
        self.writer.output_dir.write_text("existing")
        self.assertFalse(self.writer.create_file("a.md", BODY)["success"])
        self.assertEqual(self.writer.output_dir.read_text(), "existing")

    def test_output_directory_cannot_be_redirected(self):
        outside = self.project / "outside"
        outside.mkdir()
        if os.name == "nt":
            subprocess.run(["cmd", "/c", "mklink", "/J", str(self.writer.output_dir), str(outside)],
                           check=True, capture_output=True)
        else:
            self.writer.output_dir.symlink_to(outside, target_is_directory=True)
        try:
            self.assertEqual(self.writer.create_file("a.md", BODY)["code"], "unsafe_directory")
            self.assertEqual(list(outside.iterdir()), [])
        finally:
            if os.name == "nt":
                self.writer.output_dir.rmdir()
            else:
                self.writer.output_dir.unlink()

    def test_missing_pandoc_and_pdf_engine(self):
        with patch("agent.file_writer.shutil.which", return_value=None):
            self.assertEqual(self.writer.create_file("a.docx", BODY)["code"], "pandoc_missing")
        with patch("agent.file_writer.shutil.which", side_effect=lambda name: "pandoc" if name == "pandoc" else None):
            self.assertEqual(self.writer.create_file("a.pdf", BODY)["code"], "pdf_backend_missing")
        self.assertEqual(list(self.writer.output_dir.iterdir()), [])

    def test_conversion_failure_timeout_and_chinese_font_failure(self):
        backend = PandocBackend()
        for result, code in [(subprocess.CompletedProcess([], 1, "", "fontspec error: font not found"), "conversion_failed"),
                             (subprocess.CompletedProcess([], 0, "", "Missing character: There is no 中"), "missing_glyph")]:
            with patch("agent.file_writer.subprocess.run", return_value=result):
                with self.assertRaises(FileWriteError) as error:
                    backend._run(["pandoc"], self.project, "body")
                self.assertEqual(error.exception.code, code)
        with patch("agent.file_writer.subprocess.run", side_effect=subprocess.TimeoutExpired("pandoc", 90)):
            with self.assertRaises(FileWriteError) as error:
                backend._run(["pandoc"], self.project, "body")
            self.assertEqual(error.exception.code, "conversion_timeout")
        with patch.object(self.writer.pandoc, "convert", side_effect=FileWriteError("conversion_failed", "失败")):
            self.assertFalse(self.writer.create_file("a.pdf", BODY)["success"])
        self.assertEqual(list(self.writer.output_dir.iterdir()), [])

    def test_empty_output_rejected(self):
        with patch.object(self.writer.pandoc, "convert", return_value=None):
            self.assertEqual(self.writer.create_file("a.docx", BODY)["code"], "empty_output")
        self.assertEqual(list(self.writer.output_dir.iterdir()), [])

    def test_images_not_loaded_and_raw_tex_not_executed(self):
        self.assertEqual(self.writer.create_file("a.docx", "![image](https://example.com/a.png)")["code"], "unsupported_content")
        body = r"$\input{secret.txt}$" + "\n\n" + r"\input{secret.txt}"
        (self.project / "secret.txt").write_text("SECRET_SHOULD_NOT_APPEAR")
        result = self.writer.create_file("safe.docx", body)
        self.assertTrue(result["success"], result)
        from docx import Document
        text = "\n".join(p.text for p in Document(result["path"]).paragraphs)
        self.assertNotIn("SECRET_SHOULD_NOT_APPEAR", text)
        self.assertIn("input", text)

    def test_tool_schema_runtime_and_memory_separation(self):
        spec = next(s for s in TOOLS if s["name"] == "create_file")
        self.assertEqual(set(spec["input_schema"]["properties"]), {"filename", "content"})
        memory = self.project / "memory"
        memory.mkdir()
        (memory / "AGENT.md").write_text("Use evidence.")
        files = FileTools(memory)
        for flag in ("read_only", "processing_eml", "edit_learning"):
            setattr(files, flag, True)
            self.assertFalse(files.execute("create_file", {"filename": "a.md", "content": BODY})["success"])
            setattr(files, flag, False)
        self.assertFalse(files.execute("create_file", {"path": "a.md", "content": BODY})["success"])
        responses = [dict(content=[dict(type="tool_use", id="1", name="create_file", input={"filename": "a.md", "content": BODY})], stop_reason="tool_use"),
                     dict(content=[dict(type="tool_use", id="2", name="create_file", input={"filename": "a.md", "content": BODY})], stop_reason="tool_use"),
                     dict(content=[dict(type="text", text="已生成，第 2 次同名请求被拒绝")], stop_reason="end_turn")]
        client = Mock()
        client.complete.side_effect = responses
        messages = []
        run_turn(client, files, messages, "保存为 a.md", emit=lambda _: None)
        self.assertTrue(json.loads(messages[2]["content"][0]["content"])["success"])
        self.assertTrue(messages[4]["content"][0]["is_error"])
        self.assertFalse(files.policy.changes)
        self.assertFalse((memory / "a.md").exists())


if __name__ == "__main__":
    unittest.main()
